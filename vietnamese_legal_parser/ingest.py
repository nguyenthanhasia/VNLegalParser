from __future__ import annotations

import json
from pathlib import Path
from typing import Any, Iterable

from .normalization import html_to_text


def read_text_file(path: str | Path) -> str:
    p = Path(path)
    suffix = p.suffix.lower()
    if suffix in {".txt", ".md"}:
        return p.read_text(encoding="utf-8", errors="replace")
    if suffix in {".html", ".htm"}:
        return html_to_text(p.read_text(encoding="utf-8", errors="replace"))
    if suffix == ".pdf":
        try:
            from pypdf import PdfReader  # type: ignore
        except ImportError as exc:  # pragma: no cover - optional dependency
            raise RuntimeError("PDF support requires: pip install 'vietnamese-legal-parser[pdf]'") from exc
        reader = PdfReader(str(p))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    if suffix == ".docx":
        try:
            from docx import Document  # type: ignore
            from docx.table import Table  # type: ignore
            from docx.text.paragraph import Paragraph  # type: ignore
        except ImportError as exc:  # pragma: no cover - optional dependency
            raise RuntimeError("DOCX support requires: pip install 'vietnamese-legal-parser[docx]'") from exc

        doc = Document(str(p))
        parts: list[str] = []
        # python-docx >=1.1 exposes document-order paragraphs/tables via iter_inner_content().
        # Keep a compatibility fallback for older installations despite our declared minimum.
        if hasattr(doc, "iter_inner_content"):
            blocks = doc.iter_inner_content()
        else:  # pragma: no cover - compatibility path
            blocks = list(doc.paragraphs) + list(doc.tables)
        for block in blocks:
            if isinstance(block, Paragraph):
                if block.text.strip():
                    parts.append(block.text)
            elif isinstance(block, Table):
                for row in block.rows:
                    cells: list[str] = []
                    previous = None
                    for cell in row.cells:
                        value = " ".join(p.text.strip() for p in cell.paragraphs if p.text.strip()).strip()
                        # Merged cells can be repeated by python-docx; avoid duplicate text.
                        if value and value != previous:
                            cells.append(value)
                        previous = value
                    if cells:
                        parts.append("[TABLE] " + " | ".join(cells))
        return "\n".join(parts)
    raise ValueError(f"Unsupported input format: {suffix}")


def load_records(path: str | Path) -> Iterable[dict[str, Any]]:
    p = Path(path)
    suffix = p.suffix.lower()
    if suffix == ".jsonl":
        with p.open("r", encoding="utf-8") as f:
            for line_no, line in enumerate(f, 1):
                if not line.strip():
                    continue
                try:
                    obj = json.loads(line)
                except json.JSONDecodeError as exc:
                    raise ValueError(f"Invalid JSONL at line {line_no}: {exc.msg}") from exc
                if not isinstance(obj, dict):
                    raise ValueError(f"JSONL line {line_no} is not an object")
                yield obj
        return
    if suffix == ".json":
        obj = json.loads(p.read_text(encoding="utf-8"))
        if isinstance(obj, list):
            for item in obj:
                if not isinstance(item, dict):
                    raise ValueError("JSON array items must be objects")
                yield item
        elif isinstance(obj, dict):
            yield obj
        else:
            raise ValueError("JSON must be an object or an array of objects")
        return
    yield {"document_id": p.stem, "title": p.stem, "content": read_text_file(p), "source": str(p)}
